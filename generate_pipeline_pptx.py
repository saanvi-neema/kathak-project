"""
Generates kathak_pipeline.pptx — an editable PowerPoint version of the
pipeline flow diagram. All nodes are shapes, all arrows are connectors.
Edit freely in PowerPoint: drag, recolor, add slides, etc.
"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_AUTO_SIZE, MSO_ANCHOR
from pptx.oxml.ns import qn
from lxml import etree

# ---------------------------------------------------------------------------
# Color palette (matches pipeline_flow.html CSS variables)
# ---------------------------------------------------------------------------
def rgb(r, g, b): return RGBColor(r, g, b)

TEAL        = rgb(0x1B, 0x6E, 0x67)
TEAL_SOFT   = rgb(0xE4, 0xEF, 0xEA)
GOLD        = rgb(0xA8, 0x75, 0x2B)
GOLD_SOFT   = rgb(0xF3, 0xE7, 0xCE)
CORAL       = rgb(0xC0, 0x39, 0x2B)
CORAL_SOFT  = rgb(0xFD, 0xEA, 0xEA)
SURFACE2    = rgb(0xFB, 0xF6, 0xEA)
RULE        = rgb(0xDE, 0xD0, 0xAF)
TERMINAL    = rgb(0xFF, 0xFD, 0xF6)
TERM_RULE   = rgb(0xC9, 0xB7, 0x7E)
INK         = rgb(0x21, 0x1C, 0x14)
INK_DIM     = rgb(0x6E, 0x66, 0x56)
BG          = rgb(0xF6, 0xF1, 0xE6)

# ---------------------------------------------------------------------------
# Coordinate system: SVG → slide inches
# SVG logical bounds: x 40–1280, y 68–600
# Slide: 13.33" × 7.5" (16:9 widescreen)
# ---------------------------------------------------------------------------
SVG_X0, SVG_Y0 = 40, 68
SVG_W, SVG_H   = 1240, 532   # 1280-40, 600-68
SLIDE_CW = 12.5              # usable width
SLIDE_CH = 6.5               # usable height
XM, YM   = 0.4, 0.5         # left/top margin in inches

_xs = SLIDE_CW / SVG_W
_ys = SLIDE_CH / SVG_H

def sx(v):  return Inches((v - SVG_X0) * _xs + XM)
def sy(v):    return Inches((v - SVG_Y0) * _ys + YM)
def sw(v):    return Inches(v * _xs)
def sh_in(v): return Inches(v * _ys)

# Raw float helpers (for connector points)
def fx(v):  return (v - SVG_X0) * _xs + XM
def fy(v):  return (v - SVG_Y0) * _ys + YM

# ---------------------------------------------------------------------------
# Shape helpers
# ---------------------------------------------------------------------------
ROUNDED_RECT = 5   # MSO autoshape id for rounded rectangle

def add_box(shapes, svg_x, svg_y, svg_w, svg_h,
            fill, border, rows,
            border_w_pt=1.2, corner_round=0.05):
    """Add a rounded rectangle with multi-row text."""
    shape = shapes.add_shape(
        ROUNDED_RECT,
        sx(svg_x), sy(svg_y), sw(svg_w), sh_in(svg_h)
    )
    # Corner rounding: set adj XML attribute
    try:
        prstGeom = shape._element.spPr.find(qn('a:prstGeom'))
        if prstGeom is not None:
            avLst = prstGeom.find(qn('a:avLst'))
            if avLst is None:
                avLst = etree.SubElement(prstGeom, qn('a:avLst'))
            gd = avLst.find(qn('a:gd'))
            if gd is None:
                gd = etree.SubElement(avLst, qn('a:gd'))
            gd.set('name', 'adj')
            gd.set('fmla', f'val {int(corner_round * 50000)}')
    except Exception:
        pass

    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = border
    shape.line.width = Pt(border_w_pt)

    tf = shape.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    # Vertical center via XML bodyPr anchor
    txBody = tf._txBody
    bodyPr = txBody.find(qn('a:bodyPr'))
    if bodyPr is not None:
        bodyPr.set('anchor', 'ctr')
    bodyPr.set('lIns', str(Emu(Inches(0.08))))
    bodyPr.set('rIns', str(Emu(Inches(0.08))))

    for i, row in enumerate(rows):
        text, size, bold, color = row
        if i == 0:
            para = tf.paragraphs[0]
        else:
            para = tf.add_paragraph()
        para.alignment = PP_ALIGN.CENTER
        para.space_before = Pt(1)
        run = para.add_run()
        run.text = text
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color

    return shape


def add_connector(shapes, x1, y1, x2, y2,
                  color, width_pt=1.8, arrow_end=True,
                  dashed=False, elbow=False):
    """Add a straight or elbow connector with optional arrowhead."""
    from pptx.enum.shapes import MSO_CONNECTOR_TYPE
    ctype = MSO_CONNECTOR_TYPE.ELBOW if elbow else MSO_CONNECTOR_TYPE.STRAIGHT
    conn = shapes.add_connector(ctype,
        Inches(x1), Inches(y1), Inches(x2), Inches(y2))

    ln = conn.line
    ln.color.rgb = color
    ln.width = Pt(width_pt)

    # Arrow and dash via XML
    sp_ln = conn._element.spPr.find(qn('a:ln'))
    if sp_ln is None:
        sp_ln = etree.SubElement(conn._element.spPr, qn('a:ln'))

    if dashed:
        prstDash = etree.SubElement(sp_ln, qn('a:prstDash'))
        prstDash.set('val', 'dash')

    if arrow_end:
        headEnd = etree.SubElement(sp_ln, qn('a:headEnd'))
        headEnd.set('type', 'triangle')
        headEnd.set('w', 'med')
        headEnd.set('len', 'med')

    return conn


def add_label(shapes, x, y, w_in, text, size, color, bold=False):
    """Floating text label (for arrow annotations)."""
    txb = shapes.add_textbox(Inches(x), Inches(y), Inches(w_in), Inches(0.22))
    tf = txb.text_frame
    run = tf.paragraphs[0].add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.bold = bold
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    return txb


# ---------------------------------------------------------------------------
# Build slide
# ---------------------------------------------------------------------------
prs = Presentation()
prs.slide_width  = Inches(13.33)
prs.slide_height = Inches(7.5)

slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank layout

# Background colour
bg = slide.background
bgfill = bg.fill
bgfill.solid()
bgfill.fore_color.rgb = BG

sh = slide.shapes   # shorthand

# ---------------------------------------------------------------------------
# TITLE
# ---------------------------------------------------------------------------
tb = sh.add_textbox(Inches(0.4), Inches(0.08), Inches(12.5), Inches(0.38))
tf = tb.text_frame
p = tf.paragraphs[0]
run = p.add_run()
run.text = "Kathak Analysis Pipeline"
run.font.size = Pt(20)
run.font.bold = True
run.font.color.rgb = INK
p.alignment = PP_ALIGN.CENTER

# ---------------------------------------------------------------------------
# NODES
# ---------------------------------------------------------------------------

# --- Upload Video (teal) ---
add_box(sh, 40, 68, 190, 64, TEAL_SOFT, TEAL, [
    ("Upload Video", 14, True,  INK),
])

# --- webcam chunk / Live (gold) ---
add_box(sh, 40, 520, 190, 64, GOLD_SOFT, GOLD, [
    ("webcam chunk", 14, True,  INK),
])

# --- MediaPipe Models (neutral, dashed feel) ---
add_box(sh, 360, 190, 200, 52, SURFACE2, RULE, [
    ("MediaPipe Models",          12.5, True,  INK),
    ("pose · hand · face  .task", 10,   False, INK_DIM),
], border_w_pt=1.0)

# --- Mudra Classifier (neutral) ---
add_box(sh, 690, 130, 220, 52, SURFACE2, RULE, [
    ("Mudra Classifier",          12.5, True,  INK),
    ("trained ML model  ·  ~174 MB", 10, False, INK_DIM),
], border_w_pt=1.0)

# --- Extract landmarks + audio (Ingest) ---
add_box(sh, 350, 300, 220, 90, SURFACE2, RULE, [
    ("Extract landmarks",                    14,   True,  INK),
    ("+ audio",                              14,   True,  INK),
    ("MediaPipe pose/hand/face, librosa",    11,   False, INK_DIM),
])

# --- Analysis ---
add_box(sh, 650, 210, 310, 190, SURFACE2, RULE, [
    ("Analysis",                    15,   True,  INK),
    ("same code, both paths",       11,   False, INK_DIM),
    ("Chakkar   Mudra   Rasa",      12,   False, INK),
    ("Timing & Taal   Tatkaar",     12,   False, INK),
    ("whatever data exists so far", 11,   False, INK_DIM),
])

# --- Scoring (terminal / shared) ---
add_box(sh, 1010, 240, 250, 80, TERMINAL, TERM_RULE, [
    ("Scoring",                              15,   True,  INK),
    ("mudra_scoring() · chakkar · timing",   11,   False, INK_DIM),
    ("identical result, either path",        11,   False, INK_DIM),
])

# --- Report → Dashboard (terminal) ---
add_box(sh, 1010, 348, 250, 72, TERMINAL, TERM_RULE, [
    ("Report",       15,  True, INK),
    ("→ Dashboard",  15,  True, INK),
])

# --- Finalize-margin gate (gold / live-only) ---
add_box(sh, 1000, 460, 240, 76, GOLD_SOFT, GOLD, [
    ("Finalize-margin gate",              14,  True,  INK),
    ("only emit safely-finished events",  11,  False, INK_DIM),
    ("confirmed past · not ongoing",      10,  False, GOLD),
])

# --- Threshold gate (coral / hardware) ---
add_box(sh, 630, 432, 350, 60, CORAL_SOFT, CORAL, [
    ("Threshold gate",                          13.5, True,  INK),
    ("buzz only if score < threshold (e.g. 20%)", 11, False, INK_DIM),
    ("fires during hold · not after it ends",    10,  False, CORAL),
])

# --- Hardware / Haptic Glove (coral) ---
add_box(sh, 640, 510, 330, 68, CORAL_SOFT, CORAL, [
    ("Haptic Glove",                       14,  True,  INK),
    ("Arduino · vibration motor per finger", 12, False, INK_DIM),
    ("buzz finger with low mudra score",    11,  False, CORAL),
])

# ---------------------------------------------------------------------------
# ARROWS
# Notation: (x1,y1) → (x2,y2) in SVG px, converted to inches via fx/fy
# ---------------------------------------------------------------------------

# a: Upload bottom-center → Ingest top-center (teal)
add_connector(sh,
    fx(135), fy(132),   # Upload bottom
    fx(460), fy(300),   # Ingest top
    TEAL, width_pt=2.0, elbow=True)
add_label(sh, fx(135)-0.3, fy(210), 1.1, "upload path", 8.5, TEAL)

# b: Live right-center → Ingest bottom-center (gold, loop back)
add_connector(sh,
    fx(230), fy(552),   # Live right
    fx(460), fy(390),   # Ingest bottom
    GOLD, width_pt=2.0, elbow=True)
add_label(sh, fx(270), fy(450), 1.1, "live path", 8.5, GOLD)

# MediaPipe Models → Ingest (neutral dashed)
add_connector(sh,
    fx(460), fy(242),   # Models bottom
    fx(460), fy(300),   # Ingest top
    RULE, width_pt=1.4, dashed=True)

# Mudra Classifier → Analysis (neutral dashed)
add_connector(sh,
    fx(800), fy(182),   # Classifier bottom
    fx(800), fy(210),   # Analysis top
    RULE, width_pt=1.4, dashed=True)

# c: Ingest right → Analysis left (neutral)
add_connector(sh,
    fx(570), fy(345),
    fx(650), fy(305),
    INK_DIM, width_pt=1.8, elbow=True)

# d: Analysis right → Scoring left (teal, upload path)
add_connector(sh,
    fx(960), fy(260),
    fx(1010), fy(280),
    TEAL, width_pt=2.0)
add_label(sh, fx(960)+0.05, fy(248), 1.2, "upload path", 8.5, TEAL)

# e: Analysis right → Finalize gate left (gold, live path)
add_connector(sh,
    fx(960), fy(370),
    fx(1000), fy(498),
    GOLD, width_pt=2.0, elbow=True)
add_label(sh, fx(960)+0.05, fy(400), 1.2, "live path", 8.5, GOLD)

# f: Finalize gate right → Scoring bottom (gold loop up)
add_connector(sh,
    fx(1240), fy(498),
    fx(1135), fy(320),
    GOLD, width_pt=2.0, elbow=True)

# k: Scoring → Report+Dashboard (down)
add_connector(sh,
    fx(1135), fy(320),
    fx(1135), fy(348),
    INK_DIM, width_pt=1.8)

# l: Scoring bottom-left → Threshold gate right (coral, score feed)
add_connector(sh,
    fx(1010), fy(320),
    fx(980),  fy(432),
    CORAL, width_pt=1.8, elbow=True)
add_label(sh, fx(985)-0.6, fy(370), 1.2, "score (live only)", 8.5, CORAL)

# h: Analysis bottom → Threshold gate top (coral, ongoing finger state)
add_connector(sh,
    fx(805), fy(402),
    fx(805), fy(432),
    CORAL, width_pt=1.8)
add_label(sh, fx(720), fy(406), 1.8, "current finger state · ongoing holds", 8.0, CORAL)

# j: Threshold gate → Hardware node
add_connector(sh,
    fx(805), fy(492),
    fx(805), fy(510),
    CORAL, width_pt=1.8)

# i: Hardware bottom → Live (loop back, coral)
add_connector(sh,
    fx(640), fy(578),
    fx(190), fy(578),
    CORAL, width_pt=1.8, elbow=True)
add_label(sh, fx(380), fy(580), 1.8, "corrected posture → retry", 8.5, CORAL)

# Hardware loop left end → Live right
add_connector(sh,
    fx(190), fy(578),
    fx(230), fy(552),
    CORAL, width_pt=1.8, elbow=True)

# ---------------------------------------------------------------------------
# LEGEND
# ---------------------------------------------------------------------------
legend_x, legend_y = 0.38, 7.08
for color, label in [(TEAL, "Upload path"), (GOLD, "Live (camera) path"), (CORAL, "Hardware (haptic) path")]:
    box = sh.add_shape(1, Inches(legend_x), Inches(legend_y), Inches(0.18), Inches(0.18))
    box.fill.solid()
    box.fill.fore_color.rgb = color
    box.line.color.rgb = color
    tb2 = sh.add_textbox(Inches(legend_x + 0.22), Inches(legend_y - 0.01), Inches(1.6), Inches(0.22))
    run = tb2.text_frame.paragraphs[0].add_run()
    run.text = label
    run.font.size = Pt(9)
    run.font.color.rgb = color
    legend_x += 2.0

# Dashed legend entry
line_shape = sh.add_shape(1, Inches(legend_x), Inches(legend_y + 0.07), Inches(0.18), Inches(0.04))
line_shape.fill.solid()
line_shape.fill.fore_color.rgb = RULE
line_shape.line.color.rgb = RULE
tb3 = sh.add_textbox(Inches(legend_x + 0.22), Inches(legend_y - 0.01), Inches(2.0), Inches(0.22))
run3 = tb3.text_frame.paragraphs[0].add_run()
run3.text = "Model dependency"
run3.font.size = Pt(9)
run3.font.color.rgb = INK_DIM

# ---------------------------------------------------------------------------
# Save
# ---------------------------------------------------------------------------
out = "kathak_pipeline.pptx"
prs.save(out)
print(f"Saved: {out}")
